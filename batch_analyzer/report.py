import random
from collections import Counter

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from classifier import OTHER

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
SENTIMENTS = ["良い", "悪い", "どちらでもない"]


def _write_table(ws, header: list[str], rows: list[list], widths: list[int], start_row: int = 1):
    for c, h in enumerate(header, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for r, row in enumerate(rows, start=start_row + 1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def category_order(rows: list[dict], categories: list[str]) -> list[str]:
    """件数の多い順。「その他」は最後。"""
    counts = Counter(r["category"] for r in rows if r["category"])
    names = [c for c in categories if c != OTHER]
    return sorted(names, key=lambda c: -counts[c]) + [OTHER]


def spot_check_indices(rows: list[dict], ratio: float = 0.1, minimum: int = 10, seed: int = 0) -> list[int]:
    """抜き取り確認用に、分類できた回答から1割（最低10件）をランダムに選ぶ。seed を固定して毎回同じものを選ぶ。"""
    ok = [i for i, r in enumerate(rows) if not r["error"]]
    k = min(len(ok), max(minimum, round(len(ok) * ratio)))
    return sorted(random.Random(seed).sample(ok, k))


def build_workbook(rows: list[dict], extra_columns: list[str], extras: list[dict], categories: list[str],
                   representatives: dict[str, list[int]], group_by: str | None = None,
                   with_details: bool = True) -> Workbook:
    """納品用のExcelを作る。

    rows: classify() の結果。extras: 元のファイルの他の列（rows と同じ順番）
    with_details: False のとき（ベーシック）は、感情・要約の列と代表的な声のシートを出さない
    """
    wb = Workbook()
    order = category_order(rows, categories)
    classified = [r for r in rows if not r["error"]]
    total = len(classified)

    # 集計シート（最初に開くシート）
    ws = wb.active
    ws.title = "集計"
    counts = Counter(r["category"] for r in classified)
    table = [[c, counts[c], counts[c] / total if total else 0] for c in order]
    ws["A1"] = f"分類ごとの件数（分類できた回答 {total}件）"
    ws["A1"].font = Font(bold=True, size=13)
    _write_table(ws, ["分類", "件数", "割合"], table, [22, 10, 10], start_row=3)
    for r in range(4, 4 + len(table)):
        ws.cell(row=r, column=3).number_format = "0.0%"
    failed = len(rows) - total
    if failed:
        ws.cell(row=5 + len(table), column=1, value=f"※ 分類できなかった回答が{failed}件あります（一覧シートの「エラー」列を参照）")

    chart = BarChart()
    chart.type = "bar"
    chart.title = "分類ごとの件数"
    chart.legend = None
    chart.add_data(Reference(ws, min_col=2, min_row=3, max_row=3 + len(table)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + len(table)))
    chart.x_axis.scaling.orientation = "maxMin"  # 件数の多い分類を上に
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "E3")

    if with_details:
        row0 = 6 + len(table) + (1 if failed else 0)
        ws.cell(row=row0, column=1, value="分類×感情").font = Font(bold=True)
        cross = [[c] + [sum(1 for r in classified if r["category"] == c and r["sentiment"] == s) for s in SENTIMENTS]
                 for c in order]
        _write_table(ws, ["分類"] + SENTIMENTS, cross, [22, 10, 10, 14], start_row=row0 + 1)

    # クロス集計（例：分類×年代）
    if group_by:
        wc = wb.create_sheet(f"分類×{group_by}"[:31])
        groups = sorted({e.get(group_by, "") for e in extras})
        cross = [[c] + [sum(1 for r, e in zip(rows, extras) if not r["error"] and r["category"] == c
                            and e.get(group_by, "") == g) for g in groups] for c in order]
        wc["A1"] = f"分類×{group_by}（件数）"
        wc["A1"].font = Font(bold=True, size=13)
        _write_table(wc, ["分類"] + [g or "（空欄）" for g in groups], cross, [22] + [12] * len(groups), start_row=3)

    # 一覧シート
    wl = wb.create_sheet("一覧")
    header = ["No."] + extra_columns + ["回答", "分類"]
    widths = [6] + [14] * len(extra_columns) + [60, 16]
    if with_details:
        header += ["感情", "ひと言要約"]
        widths += [14, 36]
    header.append("エラー")
    widths.append(20)
    body = []
    for n, (r, e) in enumerate(zip(rows, extras), start=1):
        line = [n] + [e.get(c, "") for c in extra_columns] + [r["text"], r["category"]]
        if with_details:
            line += [r["sentiment"], r["summary"]]
        body.append(line + [r["error"]])
    _write_table(wl, header, body, widths)
    wl.auto_filter.ref = wl.dimensions
    wl.freeze_panes = "A2"

    # 代表的な声
    if with_details:
        wr = wb.create_sheet("代表的な声")
        voices = [[c, counts[c], rows[i]["text"]] for c in order for i in representatives.get(c, [])]
        _write_table(wr, ["分類", "件数", "回答"], voices, [18, 8, 80])
    return wb


def build_spot_check(rows: list[dict]) -> Workbook:
    """抜き取り確認用のExcel（納品前に当方が読んで確認する。お客様には渡さない）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "抜き取り確認"
    picks = spot_check_indices(rows)
    _write_table(ws, ["No.", "回答", "分類", "確認結果（○ / 正しい分類）"],
                 [[i + 1, rows[i]["text"], rows[i]["category"], ""] for i in picks], [6, 60, 16, 24])
    return wb
