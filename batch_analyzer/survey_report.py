import sys
import csv
import json
import random
import argparse
from pathlib import Path

from anthropic import APIError
from openpyxl import load_workbook

from classifier import (ClassificationError, OTHER, classify, estimate_note, pick_representatives,
                        suggest_categories)
from report import build_spot_check, build_workbook, category_order


def log(msg: str):
    print(msg, file=sys.stderr)


def read_table(path: Path) -> list[dict]:
    """CSV（UTF-8 / Excelで保存したShift_JIS）または Excel（1枚目のシート）を読み込む。"""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        values = [[("" if v is None else str(v)) for v in row] for row in ws.iter_rows(values_only=True)]
        header, body = values[0], values[1:]
        return [dict(zip(header, row)) for row in body if any(row)]
    for encoding in ("utf-8-sig", "cp932"):
        try:
            with open(path, encoding=encoding, newline="") as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"[error] {path.name} の文字コードを判定できませんでした（UTF-8 か Shift_JIS で保存してください）")


def load_categories(value: str) -> list[dict]:
    """「接客,価格,品質」のようなカンマ区切り、または suggest で作ったJSONファイルを読み込む。"""
    path = Path(value)
    if path.suffix == ".json" and path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return [{"name": name.strip(), "description": ""} for name in value.split(",") if name.strip()]


def cmd_suggest(args):
    rows = read_table(Path(args.input))
    texts = [r[args.column] for r in rows if r.get(args.column, "").strip()]
    sample = random.Random(0).sample(texts, min(args.sample, len(texts)))
    log(f"[info] {len(texts)}件のうち{len(sample)}件を読んで、分類案を作ります")
    ideas, usage = suggest_categories(sample, args.n, args.purpose)
    Path(args.output).write_text(json.dumps(ideas, ensure_ascii=False, indent=2), encoding="utf-8")
    for c in ideas:
        print(f"- {c['name']}：{c['description']}")
    log(f"[done] 分類案を {args.output} に保存しました（お客様に確認してもらい、必要なら直してください）")
    log(f"[info] {estimate_note(usage)}")


def cmd_classify(args):
    rows_in = read_table(Path(args.input))
    if not rows_in or args.column not in rows_in[0]:
        raise SystemExit(f"[error] 列「{args.column}」が見つかりません（-c で回答の列名を指定してください）")
    if args.group_by and args.group_by not in rows_in[0]:
        raise SystemExit(f"[error] 列「{args.group_by}」が見つかりません")
    categories = load_categories(args.categories)
    names = [c["name"] for c in categories if c["name"] != OTHER]
    log(f"[info] {len(rows_in)}件を {len(names)}分類＋「{OTHER}」に分けます：{'、'.join(names)}")

    texts = [r.get(args.column, "") for r in rows_in]
    rows, usage = classify(texts, categories, log=log)

    representatives = {}
    if not args.basic:
        for c in category_order(rows, names):
            try:
                representatives[c], u = pick_representatives(rows, c)
                usage = {k: usage[k] + u[k] for k in usage}
            except (ClassificationError, APIError) as e:
                log(f"[warn] 「{c}」の代表的な声を選べませんでした: {e}")

    extra_columns = [c for c in rows_in[0] if c != args.column]
    out = Path(args.output)
    build_workbook(rows, extra_columns, rows_in, names, representatives, args.group_by,
                   with_details=not args.basic).save(out)
    check = out.with_name(out.stem + "_抜き取り確認.xlsx")
    build_spot_check(rows).save(check)

    failed = sum(1 for r in rows if r["error"])
    log(f"[done] {len(rows)}件を処理（分類できた{len(rows) - failed} / 失敗{failed}）し、{out} に出力しました")
    log(f"[info] 納品前に {check} の回答を読んで、分類のずれがないか確認してください")
    log(f"[info] {estimate_note(usage)}")


def main():
    parser = argparse.ArgumentParser(description="アンケート・口コミの自由記述をAIで分類し、集計したExcelを作るツール")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("suggest", help="回答の一部を読んで、分類の案を作る（分類がおまかせのとき）")
    p.add_argument("input", help="回答のファイル（.csv / .xlsx）")
    p.add_argument("-c", "--column", default="text", help="回答の列名（デフォルト: text）")
    p.add_argument("-n", type=int, default=6, help="分類の数の目安（デフォルト: 6）")
    p.add_argument("--sample", type=int, default=60, help="読ませる回答の数（デフォルト: 60）")
    p.add_argument("--purpose", default="", help="分析の目的（例：「店舗の改善点を知りたい」）")
    p.add_argument("-o", "--output", default="categories.json", help="分類案の保存先（デフォルト: categories.json）")
    p.set_defaults(func=cmd_suggest)

    p = sub.add_parser("classify", help="全件を分類して、集計したExcelを作る")
    p.add_argument("input", help="回答のファイル（.csv / .xlsx）")
    p.add_argument("-c", "--column", default="text", help="回答の列名（デフォルト: text）")
    p.add_argument("--categories", required=True, help="「接客,価格,品質」のようなカンマ区切り、または suggest で作ったJSON")
    p.add_argument("--group-by", help="クロス集計に使う列名（例：年代）")
    p.add_argument("--basic", action="store_true", help="分類と集計だけにする（感情・要約・代表的な声を出さない）")
    p.add_argument("-o", "--output", default="survey_report.xlsx", help="出力するExcel（デフォルト: survey_report.xlsx）")
    p.set_defaults(func=cmd_classify)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
